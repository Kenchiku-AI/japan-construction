#!/usr/bin/env python3

import json
import math
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

try:
  import xlrd
except ImportError:
  xlrd = None


# ---------------------------------------------------------------------------
# General helpers
# ---------------------------------------------------------------------------

MAX_CELL_OUTPUT = 10000
MAX_STRING_LENGTH = 2000


def clean_value(value: Any) -> Any:
  """Convert Excel/Python values into JSON-safe values."""

  if value is None:
    return None

  if isinstance(value, float):
    if math.isnan(value) or math.isinf(value):
      return None

  if isinstance(value, (str, int, float, bool)):
    if isinstance(value, str) and len(value) > MAX_STRING_LENGTH:
      return value[:MAX_STRING_LENGTH] + "...[truncated]"
    return value

  if hasattr(value, "isoformat"):
    try:
      return value.isoformat()
    except Exception:
      pass

  return str(value)


def color_to_dict(color: Any) -> Any:
  """Serialize an openpyxl Color object."""

  if color is None:
    return None

  result = {}

  for attribute in (
    "type",
    "rgb",
    "indexed",
    "auto",
    "theme",
    "tint",
  ):
    value = getattr(color, attribute, None)

    if value is not None:
      result[attribute] = value

  return result or None


def side_to_dict(side: Any) -> Any:
  """Serialize an openpyxl border side."""

  if side is None:
    return None

  result = {
    "style": side.style,
  }

  color = color_to_dict(side.color)

  if color:
    result["color"] = color

  return result


def font_to_dict(font: Any) -> dict:
  return {
    "name": font.name,
    "size": font.sz,
    "bold": font.bold,
    "italic": font.italic,
    "underline": font.underline,
    "strike": font.strike,
    "color": color_to_dict(font.color),
  }


def fill_to_dict(fill: Any) -> dict:
  return {
    "type": fill.fill_type,
    "pattern": fill.patternType,
    "fg_color": color_to_dict(fill.fgColor),
    "bg_color": color_to_dict(fill.bgColor),
  }


def border_to_dict(border: Any) -> dict:
  return {
    "left": side_to_dict(border.left),
    "right": side_to_dict(border.right),
    "top": side_to_dict(border.top),
    "bottom": side_to_dict(border.bottom),
    "diagonal": side_to_dict(border.diagonal),
  }


def alignment_to_dict(alignment: Any) -> dict:
  return {
    "horizontal": alignment.horizontal,
    "vertical": alignment.vertical,
    "wrap_text": alignment.wrap_text,
    "shrink_to_fit": alignment.shrink_to_fit,
    "text_rotation": alignment.textRotation,
    "indent": alignment.indent,
  }


def protection_to_dict(protection: Any) -> dict:
  return {
    "locked": protection.locked,
    "hidden": protection.hidden,
  }


def dimension_is_non_default(dimension: Any) -> bool:
  """Determine whether a row/column dimension has meaningful customization."""

  if dimension is None:
    return False

  return any(
    (
      dimension.hidden,
      dimension.width is not None,
      dimension.height is not None,
      dimension.outlineLevel,
      dimension.collapsed,
    )
  )


def safe_range_boundaries(range_string: str):
  try:
    return range_boundaries(range_string)
  except Exception:
    return None


def range_contains_cell(range_string: str, coordinate: str) -> bool:
  boundaries = safe_range_boundaries(range_string)

  if boundaries is None:
    return False

  min_col, min_row, max_col, max_row = boundaries

  try:
    cell_col = ord(coordinate[0].upper()) - ord("A") + 1
  except Exception:
    return False

  row_digits = ""

  for character in coordinate[1:]:
    if character.isdigit():
      row_digits += character

  if not row_digits:
    return False

  cell_row = int(row_digits)

  return (
    min_col <= cell_col <= max_col
    and min_row <= cell_row <= max_row
  )


# ---------------------------------------------------------------------------
# Cell inspection
# ---------------------------------------------------------------------------

def inspect_cell(cell: Any) -> dict:
  result = {
    "coordinate": cell.coordinate,
    "value": clean_value(cell.value),
    "data_type": cell.data_type,
  }

  if cell.data_type == "f":
    result["formula"] = str(cell.value)

  if cell.number_format:
    result["number_format"] = cell.number_format

  result["font"] = font_to_dict(cell.font)
  result["fill"] = fill_to_dict(cell.fill)
  result["border"] = border_to_dict(cell.border)
  result["alignment"] = alignment_to_dict(cell.alignment)
  result["protection"] = protection_to_dict(cell.protection)

  if cell.hyperlink:
    result["hyperlink"] = {
      "target": cell.hyperlink.target,
      "location": cell.hyperlink.location,
      "tooltip": cell.hyperlink.tooltip,
    }

  if cell.comment:
    result["comment"] = {
      "text": cell.comment.text[:MAX_STRING_LENGTH],
      "author": cell.comment.author,
    }

  return result


def cell_has_meaningful_formatting(cell: Any) -> bool:
  """Return True when a blank cell still appears intentionally formatted."""

  if cell.value is not None:
    return True

  if cell.style_id != 0:
    return True

  if cell.number_format != "General":
    return True

  if cell.font != cell.parent.parent._fonts[0]:
    return True

  if cell.fill.fill_type is not None:
    return True

  border = cell.border

  if any(
    side.style is not None
    for side in (
      border.left,
      border.right,
      border.top,
      border.bottom,
    )
  ):
    return True

  alignment = cell.alignment

  if any(
    (
      alignment.horizontal is not None,
      alignment.vertical is not None,
      alignment.wrap_text is not None,
      alignment.shrink_to_fit is not None,
      alignment.textRotation != 0,
      alignment.indent != 0,
    )
  ):
    return True

  return False


# ---------------------------------------------------------------------------
# Merged cells
# ---------------------------------------------------------------------------

def inspect_merged_cells(worksheet: Any) -> list:
  merged = []

  for merged_range in worksheet.merged_cells.ranges:
    merged.append({
      "range": str(merged_range),
    })

  return sorted(
    merged,
    key=lambda item: item["range"],
  )


# ---------------------------------------------------------------------------
# Data validation
# ---------------------------------------------------------------------------

def parse_validation_formula(formula: Any) -> Any:
  if formula is None:
    return None

  value = str(formula)

  if value.startswith('"') and value.endswith('"'):
    value = value[1:-1]

  return value


def extract_inline_list_values(formula: Any) -> list | None:
  if formula is None:
    return None

  value = str(formula)

  if not (
    value.startswith('"')
    and value.endswith('"')
  ):
    return None

  value = value[1:-1]

  return value.split(",")


def inspect_data_validations(worksheet: Any) -> list:
  validations = []

  if not worksheet.data_validations:
    return validations

  for validation in worksheet.data_validations.dataValidation:
    entry = {
      "type": validation.type,
      "operator": validation.operator,
      "allow_blank": validation.allow_blank,
      "error": validation.error,
      "error_title": validation.errorTitle,
      "prompt": validation.prompt,
      "prompt_title": validation.promptTitle,
      "ranges": str(validation.sqref),
    }

    formula1 = parse_validation_formula(validation.formula1)
    formula2 = parse_validation_formula(validation.formula2)

    if formula1 is not None:
      entry["formula1"] = formula1

    if formula2 is not None:
      entry["formula2"] = formula2

    if validation.type == "list":
      inline_values = extract_inline_list_values(validation.formula1)

      if inline_values is not None:
        entry["allowed_values"] = inline_values

    validations.append(entry)

  return validations


# ---------------------------------------------------------------------------
# Row / column dimensions
# ---------------------------------------------------------------------------

def inspect_column_dimensions(worksheet: Any) -> list:
  result = []

  for key, dimension in worksheet.column_dimensions.items():
    if not dimension_is_non_default(dimension):
      continue

    result.append({
      "column": key,
      "width": dimension.width,
      "hidden": dimension.hidden,
      "outline_level": dimension.outlineLevel,
      "collapsed": dimension.collapsed,
      "min": dimension.min,
      "max": dimension.max,
    })

  return result


def inspect_row_dimensions(worksheet: Any) -> list:
  result = []

  for key, dimension in worksheet.row_dimensions.items():
    if not dimension_is_non_default(dimension):
      continue

    result.append({
      "row": key,
      "height": dimension.height,
      "hidden": dimension.hidden,
      "outline_level": dimension.outlineLevel,
      "collapsed": dimension.collapsed,
    })

  return result


# ---------------------------------------------------------------------------
# Sheet inspection
# ---------------------------------------------------------------------------

def inspect_worksheet(worksheet: Any) -> dict:
  result = {
    "title": worksheet.title,
    "state": worksheet.sheet_state,
    "max_row": worksheet.max_row,
    "max_column": worksheet.max_column,
    "dimensions": worksheet.calculate_dimension(),
    "freeze_panes": (
      str(worksheet.freeze_panes)
      if worksheet.freeze_panes
      else None
    ),
    "merged_cells": inspect_merged_cells(worksheet),
    "data_validations": inspect_data_validations(worksheet),
    "column_dimensions": inspect_column_dimensions(worksheet),
    "row_dimensions": inspect_row_dimensions(worksheet),
    "cells": [],
    "formatted_blank_cells": [],
  }

  cell_count = 0
  formatted_blank_count = 0

  for row in worksheet.iter_rows():
    for cell in row:
      if cell.value is not None:
        if cell_count < MAX_CELL_OUTPUT:
          result["cells"].append(
            inspect_cell(cell)
          )

        cell_count += 1

      elif cell_has_meaningful_formatting(cell):
        if formatted_blank_count < MAX_CELL_OUTPUT:
          result["formatted_blank_cells"].append(
            inspect_cell(cell)
          )

        formatted_blank_count += 1

  result["cell_count"] = cell_count
  result["formatted_blank_cell_count"] = formatted_blank_count

  if cell_count > MAX_CELL_OUTPUT:
    result["cells_truncated"] = True

  if formatted_blank_count > MAX_CELL_OUTPUT:
    result["formatted_blank_cells_truncated"] = True

  # Images / drawings are useful signals that the workbook contains
  # visual objects that are not ordinary cell values.
  try:
    result["image_count"] = len(worksheet._images)
  except Exception:
    result["image_count"] = 0

  try:
    result["chart_count"] = len(worksheet._charts)
  except Exception:
    result["chart_count"] = 0

  # Conditional formatting can affect the visual appearance of cells even
  # though the underlying cell style does not show it.
  try:
    result["conditional_formatting_count"] = len(
      worksheet.conditional_formatting
    )
  except Exception:
    result["conditional_formatting_count"] = 0

  return result


# ---------------------------------------------------------------------------
# XLSX / XLSM
# ---------------------------------------------------------------------------

def inspect_openpyxl_workbook(path: Path) -> dict:
  keep_vba = path.suffix.lower() == ".xlsm"

  workbook = load_workbook(
    filename=path,
    data_only=False,
    read_only=False,
    keep_vba=keep_vba,
  )

  result = {
    "format": path.suffix.lower().lstrip("."),
    "file_name": path.name,
    "sheet_names": workbook.sheetnames,
    "sheets": [],
  }

  try:
    result["defined_names"] = []

    for name in workbook.defined_names.values():
      result["defined_names"].append({
        "name": name.name,
        "value": name.attr_text,
        "hidden": name.hidden,
      })
  except Exception:
    result["defined_names"] = []

  try:
    result["named_styles"] = [
      style.name
      for style in workbook._named_styles
    ]
  except Exception:
    result["named_styles"] = []

  for worksheet in workbook.worksheets:
    result["sheets"].append(
      inspect_worksheet(worksheet)
    )

  return result


# ---------------------------------------------------------------------------
# XLS
# ---------------------------------------------------------------------------

def xls_cell_type_name(cell_type: int) -> str:
  if xlrd is None:
    return str(cell_type)

  mapping = {
    xlrd.XL_CELL_EMPTY: "empty",
    xlrd.XL_CELL_TEXT: "text",
    xlrd.XL_CELL_NUMBER: "number",
    xlrd.XL_CELL_DATE: "date",
    xlrd.XL_CELL_BOOLEAN: "boolean",
    xlrd.XL_CELL_ERROR: "error",
    xlrd.XL_CELL_BLANK: "blank",
  }

  return mapping.get(cell_type, str(cell_type))


def inspect_xls(path: Path) -> dict:
  if xlrd is None:
    raise RuntimeError(
      "xlrd is required to inspect .xls files."
    )

  workbook = xlrd.open_workbook(
    str(path),
    formatting_info=True,
  )

  result = {
    "format": "xls",
    "file_name": path.name,
    "sheet_names": workbook.sheet_names(),
    "sheets": [],
  }

  for sheet in workbook.sheets():
    sheet_result = {
      "title": sheet.name,
      "state": "visible",
      "max_row": sheet.nrows,
      "max_column": sheet.ncols,
      "dimensions": (
        f"A1:{get_column_letter(sheet.ncols)}{sheet.nrows}"
        if sheet.nrows and sheet.ncols
        else "A1"
      ),
      "merged_cells": [],
      "data_validations": [],
      "column_dimensions": [],
      "row_dimensions": [],
      "cells": [],
      "formatted_blank_cells": [],
      "cell_count": 0,
      "formatted_blank_cell_count": 0,
      "image_count": 0,
      "chart_count": 0,
      "conditional_formatting_count": 0,
    }

    for row_index in range(sheet.nrows):
      for column_index in range(sheet.ncols):
        cell = sheet.cell(
          row_index,
          column_index,
        )

        coordinate = (
          f"{get_column_letter(column_index + 1)}"
          f"{row_index + 1}"
        )

        value = clean_value(cell.value)

        if value is not None:
          sheet_result["cells"].append({
            "coordinate": coordinate,
            "value": value,
            "data_type": xls_cell_type_name(
              cell.ctype
            ),
          })

          sheet_result["cell_count"] += 1

    for merged_range in sheet.merged_cells:
      row_low, row_high, col_low, col_high = merged_range

      sheet_result["merged_cells"].append({
        "range": (
          f"{get_column_letter(col_low + 1)}{row_low + 1}:"
          f"{get_column_letter(col_high)}{row_high}"
        ),
      })

    result["sheets"].append(sheet_result)

  return result


# ---------------------------------------------------------------------------
# Main dispatcher
# ---------------------------------------------------------------------------

def inspect_excel(path: Path) -> dict:
  suffix = path.suffix.lower()

  if suffix in (".xlsx", ".xlsm"):
    return inspect_openpyxl_workbook(path)

  if suffix == ".xls":
    return inspect_xls(path)

  raise ValueError(
    f"Unsupported Excel format: {suffix}"
  )


def main() -> int:
  if len(sys.argv) != 2:
    print(
      f"Usage: {Path(sys.argv[0]).name} <excel-file>",
      file=sys.stderr,
    )
    return 2

  path = Path(sys.argv[1]).resolve()

  if not path.exists():
    print(
      json.dumps(
        {
          "error": f"File not found: {path}",
        },
        ensure_ascii=False,
        indent=2,
      )
    )
    return 1

  if not path.is_file():
    print(
      json.dumps(
        {
          "error": f"Not a file: {path}",
        },
        ensure_ascii=False,
        indent=2,
      )
    )
    return 1

  try:
    result = inspect_excel(path)

    print(
      json.dumps(
        result,
        ensure_ascii=False,
        indent=2,
        default=str,
      )
    )

    return 0

  except Exception as exc:
    print(
      json.dumps(
        {
          "error": (
            f"{type(exc).__name__}: {exc}"
          ),
        },
        ensure_ascii=False,
        indent=2,
      )
    )

    return 1


if __name__ == "__main__":
  raise SystemExit(main())